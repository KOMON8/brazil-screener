# -*- coding: utf-8 -*-
"""
ブラジル株スクリーナー（無料のヤフーファイナンスのデータ / yfinance 使用）

使い方:
  1. run.bat をダブルクリック（初回は必要なライブラリを自動で入れます）
  2. 同じフォルダに「ブラジル株スクリーナー.xlsx」ができます
  3. Excel の「設定」シートで条件を変えると、その場で「結果」シートが変わります
  4. 次回 run.bat を実行するとデータだけ最新になり、設定は引き継がれます
     ※ 実行する前に Excel ファイルは閉じてください

対象銘柄を変えたいときは、同じフォルダに tickers.txt を作り、1行に1銘柄
（例: PETR4）を書いてください。あればそちらを優先して使います。

注意: yfinance は非公式の方法でヤフーからデータを取っています。
      ヤフー側の変更で急に動かなくなることがあります。
"""
import datetime as dt
import json
import math
import os
import sys
import time

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import LineChart, Reference

HERE = os.path.dirname(os.path.abspath(__file__))
OUT_FILE = os.path.join(HERE, "ブラジル株スクリーナー.xlsx")
NAME_CACHE = os.path.join(HERE, "name_cache.json")
TICKER_FILE = os.path.join(HERE, "tickers.txt")
ALL_TICKER_CACHE = os.path.join(HERE, "all_tickers.json")
EARN_CACHE = os.path.join(HERE, "earnings_cache.json")
UNIVERSES = ["全銘柄", "Ibovespaの主な銘柄"]
LIQ_FOR_DETAIL = 5.0   # 平均売買代金（百万R$）がこれ未満の銘柄は、1時間足・決算日の取得を省く（時間短縮）

# Ibovespa の主な構成銘柄（2026年10月に取得確認した銘柄。入れ替えがあるので tickers.txt で調整可）
# 上場廃止・コード変更で取れない銘柄は自動でスキップし、「取得できなかった銘柄」に表示します。
DEFAULT_TICKERS = """
ALOS3 ABEV3 ASAI3 AURE3 AZZA3 B3SA3 BBSE3 BBDC3 BBDC4 BRAP4 BBAS3 BRKM5 BRAV3
BPAC11 CXSE3 CEAB3 CMIG4 COGN3 CPLE3  CSAN3 CPFE3 CMIN3 CURY3 CVCB3 CYRE3
DIRR3   AXIA3  EMBJ3 ENGI11 ENEV3 EGIE3 EQTL3 FLRY3 GGBR4 GOAU4
HAPV3 HYPE3 IGTI11 IRBR3 ISAE4 ITSA4 ITUB4 KLBN11 RENT3 LREN3 MGLU3 POMO4 
MBRF3  BEEF3 MOTV3 MRVE3 MULT3 NATU3 PETR3 PETR4 RECV3 PRIO3  PSSA3
RADL3 RAIZ4 RDOR3 RAIL3 SBSP3 SANB11  SMTO3 CSNA3 SLCE3 SUZB3 TAEE11 VIVT3
TIMS3 TOTS3 UGPA3 USIM5 VALE3 VAMO3 VBBR3 VIVA3 WEGE3 YDUQ3 SMFT3
""".split()

HISTORY_PERIOD = "5y"       # 取得する過去データの長さ（月足のMAに使うので長め）
FETCH_EARNINGS = True        # 決算予定日を取得するか（銘柄ごとに通信するので少し時間がかかる）

# 移動平均（単純移動平均）の初期設定。Excel の「設定」シートで変更でき、次回の実行から反映されます。
#  足, 期間(本), 角度を測る本数
MA_DEFAULTS = [("時間足", 25, 5), ("日足", 25, 5), ("週足", 13, 3), ("月足", 12, 3)]
TIMEFRAMES = [tf for tf, _, _ in MA_DEFAULTS]
HOURLY_PERIOD = "60d"      # 時間足は直近60日分を取得（ヤフーの1時間足は最大で約2年分）

# ---------------------------------------------------------------------------
# 条件の定義（設定シートの行になる）
#  key, 表示名, データ列の見出し, 初期ON/OFF, 初期下限, 初期上限, 説明
# ---------------------------------------------------------------------------
CONDITIONS = [
    ("budget", "投資可能額で買える", "最低購入額(R$)", "ON", None, None,
     "最低購入額（終値×売買単位）が投資可能額(R$)以下の銘柄だけ残す。下限・上限は使いません"),
    ("ang_h", "MA角度・時間足(°)", "時間足MA角度(°)", "OFF", 0, None,
     "1時間足の移動平均の角度。直近の短い流れを見る。ブラジル市場の取引時間中の足のみ"),
    ("ang_d", "MA角度・日足(°)", "日足MA角度(°)", "OFF", 0, None,
     "日足の移動平均の傾きを角度にしたもの。+は上向き、-は下向き。1本あたり1%の変化で45°"),
    ("ang_w", "MA角度・週足(°)", "週足MA角度(°)", "OFF", 0, None,
     "週足の移動平均の角度。例：緩やかな上昇は下限0・上限20 など"),
    ("ang_m", "MA角度・月足(°)", "月足MA角度(°)", "OFF", None, None,
     "月足の移動平均の角度"),
    ("chg1", "騰落率・前日比(%)", "前日比(%)", "OFF", None, None,
     "値上がり狙い→下限2・上限10 など。値下がり狙い→下限-10・上限-2 など"),
    ("chg5", "騰落率・5日(%)", "5日騰落率(%)", "OFF", None, None,
     "直近5営業日の値動き。押し目なら上限をマイナスに"),
    ("price", "株価(R$)", "終値(R$)", "OFF", 5, None,
     "低位株を外したいときに下限を設定"),
    ("value", "流動性・平均売買代金(百万R$)", "平均売買代金(百万R$)", "ON", 50, None,
     "20日平均。50なら1日5,000万レアル以上"),
    ("volr", "流動性・出来高倍率(倍)", "出来高倍率(倍)", "OFF", 1.5, None,
     "最新日の出来高 ÷ 20日平均出来高。注目度の急上昇を拾う"),
    ("atr", "ボラティリティ・ATR(%)", "ATR(%)", "ON", 2, 6,
     "14日平均の1日の値幅が株価の何%か。スイングは2〜5%、デイトレは3%以上が目安"),
    ("volup", "ボラティリティ・上昇倍率(倍)", "ボラ上昇倍率(倍)", "OFF", 1.2, None,
     "直近5日の平均値幅 ÷ 20日の平均値幅。1より大きいほど最近動きが大きくなっている"),
    ("hv", "ボラティリティ・HV20(%)", "HV20(%)", "OFF", None, None,
     "20日のヒストリカル・ボラティリティ（年率）"),
    ("dev25", "25日線からの乖離(%)", "25日線乖離(%)", "OFF", None, None,
     "終値が25日移動平均より何%上か（マイナスは下）"),
    ("rsi", "RSI(14)", "RSI(14)", "OFF", 30, 70, "押し目買いなら上限50 など"),
    ("earn", "決算までの日数", "決算までの日数", "OFF", 5, None,
     "下限5なら5日以内に決算がある銘柄を外す。決算日が不明な銘柄は残す"),
]

# プリセット（値動きのパターン）。「設定」シートで選んで使う
PATTERN_NAMES = ["なし", "急落後の落ち着き", "連騰後の失速"]
#  key, 表示名, 初期値, 説明
PATTERN_ITEMS = [
    ("sel", "使うプリセット", "なし",
     "なし／急落後の落ち着き／連騰後の失速 から選ぶ。下の条件表で ON にした条件と、両方を満たす銘柄が残ります"),
    ("mode", "条件表との組み合わせ", "投資可能額と流動性だけ",
     "プリセットを使うとき、条件表のどの条件も併用するか。「投資可能額と流動性だけ」＝投資可能額・株価・売買代金だけ併用、「すべて」＝ON の条件をすべて併用"),
    ("look", "急落・連騰を探す日数", 5, "最新日の前の何営業日を見るか。※ここを変えたら run を実行し直してください"),
    ("c_drop", "急落：1日の下落率（%以下）", -5, "この日数の中に、1日でこれ以上下げた日があれば「急落あり」"),
    ("c_lo", "落ち着き：前日比 下限（%）", -1.5, "最新日の値動きが小さい＝下げ止まり"),
    ("c_hi", "落ち着き：前日比 上限（%）", 3, ""),
    ("c_rng", "落ち着き：値幅÷ATR（以下）", 0.8, "最新日の値幅が、ふだんの値幅（ATR）の何倍か。小さいほど落ち着いた"),
    ("c_vol", "落ち着き：出来高÷急落日の出来高（以下）", 0.8, "急落した日より出来高が減っている＝売りが一巡"),
    ("c_hang", "落ち着き：時間足MA角度（°以上）", -20,
     "短い流れの下げがゆるやかになってきた（急落の1〜2日後は時間足MAがまだ下向きなので、きつくしすぎない）。空欄なら見ない"),
    ("c_rsil", "落ち着き：RSI 下限", 20, "売られすぎ圏。※急落した日の安値を割っていないことも条件に入ります"),
    ("c_rsih", "落ち着き：RSI 上限", 60, "上昇中の銘柄が急落した場合、RSIは50〜60くらいまでしか下がらないことが多い"),
    ("r_days", "連騰：連続上昇日数（以上）", 3, "前日までの連続上昇日数。これか下の5日騰落率のどちらかを満たせば「連騰あり」"),
    ("r_pct", "連騰：前日までの5日騰落率（%以上）", 8, ""),
    ("r_chg", "失速：前日比 上限（%）", 1, "最新日の伸びが小さい、またはマイナス"),
    ("r_wick", "失速：上ヒゲ比率（以上・0〜1）", 0.5, "高値を更新できなかった、または上ヒゲがこの比率以上なら「伸び悩み」"),
    ("r_vol", "失速：出来高÷直近の最大出来高（以下）", 0.8, "連騰中のピークより出来高が細ってきた"),
    ("r_hang", "失速：時間足MA角度（°以下）", 5, "短い流れが横ばい〜下向きに。空欄なら見ない"),
    ("r_rsi", "失速：RSI 下限", 65, "買われすぎ圏"),
]
PATTERN_LABELS = {label: key for key, label, *_ in PATTERN_ITEMS}
PATTERN_MODES = ["投資可能額と流動性だけ", "すべて"]
BASE_KEYS = ("budget", "price", "value")   # プリセット使用中も常に効かせる条件

DATA_COLUMNS = [
    "コード", "銘柄名", "終値(R$)", "前日比(%)", "最低購入額(R$)", "平均売買代金(百万R$)",
    "時間足MA角度(°)", "日足MA角度(°)", "週足MA角度(°)", "月足MA角度(°)",
    "時間足MA乖離(%)", "日足MA乖離(%)", "週足MA乖離(%)", "月足MA乖離(%)",
    "5日騰落率(%)", "平均出来高(株)", "出来高倍率(倍)",
    "ATR(%)", "ボラ上昇倍率(倍)", "HV20(%)", "25日線乖離(%)",
    "RSI(14)",
    "直近の最大下落率(%)", "急落からの日数", "安値更新なし", "値幅÷ATR", "出来高÷急落日",
    "連続上昇日数", "5日騰落率・前日まで(%)", "高値更新なし", "上ヒゲ比率", "出来高÷直近最大",
    "セクター", "次回決算日", "決算までの日数",
    "値幅%(4日前)", "値幅%(3日前)", "値幅%(2日前)", "値幅%(前日)", "値幅%(最新日)",
    "最終データ日", "時間足の最終時刻",
]

FONT = "Arial"
# 数式でシートを参照するときの書き方（Numbers でも読めるように引用符で囲む）
SET = "'設定'!"
DAT = "'データ'!"
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


# ---------------------------------------------------------------------------
# データ取得
# ---------------------------------------------------------------------------
def load_tickers(universe="全銘柄"):
    """対象銘柄の一覧。tickers.txt があれば最優先。全銘柄は brapi.dev の一覧（7日ごとに更新）"""
    if os.path.exists(TICKER_FILE):
        with open(TICKER_FILE, encoding="utf-8") as f:
            ts = [l.strip().upper() for l in f if l.strip() and not l.startswith("#")]
        if ts:
            return ts, {}
    if universe == "全銘柄":
        ts, info = fetch_all_tickers()
        if ts:
            return ts, info
        print("全銘柄の一覧を取得できなかったので、Ibovespaの主な銘柄で実行します。")
    return list(dict.fromkeys(DEFAULT_TICKERS)), {}


def _get_json(url):
    import urllib.request
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0", "Accept": "application/json"})
    with urllib.request.urlopen(req, timeout=30) as r:
        return json.loads(r.read().decode("utf-8"))


def fetch_all_tickers():
    """B3 の株式（普通株・優先株・ユニット）の一覧。FII・ETF・BDR は含めない"""
    import re
    pat = re.compile(r"^[A-Z0-9]{4}(3|4|5|6|7|8|11)$")
    cache = {}
    if os.path.exists(ALL_TICKER_CACHE):
        try:
            with open(ALL_TICKER_CACHE, encoding="utf-8") as f:
                cache = json.load(f)
        except Exception:
            cache = {}
    fresh = cache.get("date") and (dt.date.today() - dt.date.fromisoformat(cache["date"])).days < 7
    if fresh and cache.get("tickers"):
        return cache["tickers"], cache.get("info", {})

    print("B3の全銘柄の一覧を取得中（brapi.dev）...")
    items = []
    try:                                   # 新しい一覧API（ページ分け）
        page = 1
        while page <= 30:
            d = _get_json(f"https://brapi.dev/api/v2/tickers?type=stock&sortBy=symbol&sortOrder=asc&limit=100&page={page}")
            items += d.get("results") or []
            tp = (d.get("pagination") or {}).get("totalPages") or 1
            if page >= tp:
                break
            page += 1
            time.sleep(0.3)
    except Exception as e:
        print("  一覧API(v2)から取得できませんでした:", e)
    if not items:
        try:                               # 以前からある一覧API
            d = _get_json("https://brapi.dev/api/quote/list?type=stock&limit=2000")
            items = d.get("stocks") or []
        except Exception as e:
            print("  一覧APIから取得できませんでした:", e)

    tickers, info = [], {}
    for it in items:
        sym = str(it.get("symbol") or it.get("stock") or "").upper().strip()
        sub = str(it.get("subType") or it.get("assetType") or it.get("type") or "").lower()
        if not pat.match(sym) or sub in ("fii", "etf", "bdr", "fund", "fi-infra", "fi-agro", "fip", "fidc"):
            continue
        tickers.append(sym)
        info[sym] = {"name": it.get("name") or "", "sector": it.get("sector") or ""}
    tickers = sorted(set(tickers))
    if len(tickers) >= 100:
        print(f"  {len(tickers)} 銘柄")
        try:
            with open(ALL_TICKER_CACHE, "w", encoding="utf-8") as f:
                json.dump({"date": dt.date.today().isoformat(), "tickers": tickers, "info": info}, f, ensure_ascii=False)
        except Exception:
            pass
        return tickers, info
    if cache.get("tickers"):                # 取れなければ前回の一覧
        print("  前回の一覧を使います")
        return cache["tickers"], cache.get("info", {})
    return [], {}


def download_chunks(symbols, period, interval, chunk=100):
    """ヤフーにまとめて頼みすぎないよう、100銘柄ずつ取得"""
    import yfinance as yf
    out = {}
    for i in range(0, len(symbols), chunk):
        part = symbols[i:i + chunk]
        try:
            raw = yf.download(part, period=period, interval=interval, group_by="ticker",
                              auto_adjust=False, threads=True, progress=False)
        except Exception as e:
            print(f"  一部を取得できませんでした（{i + 1}〜{i + len(part)}）:", e)
            continue
        for s in part:
            try:
                out[s] = raw[s] if isinstance(raw.columns, pd.MultiIndex) else raw
            except Exception:
                pass
        if i + chunk < len(symbols):
            time.sleep(2)
    return out


def fetch_prices(codes):
    symbols = [c + ".SA" for c in codes]
    print(f"{len(symbols)} 銘柄の株価を取得中...")
    raw = download_chunks(symbols, HISTORY_PERIOD, "1d")
    out = {}
    for c, s in zip(codes, symbols):
        try:
            df = raw[s]
            df = clean_prices(df[["Open", "High", "Low", "Close", "Volume"]])
            if len(df) >= 30:
                out[c] = df
        except Exception:
            pass
    return out


def fetch_hourly(codes):
    """1時間足を取得（失敗しても日足の計算は続ける）"""
    symbols = [c + ".SA" for c in codes]
    print(f"{len(symbols)} 銘柄の1時間足を取得中...")
    out = {}
    raw = download_chunks(symbols, HOURLY_PERIOD, "60m")
    for c, s in zip(codes, symbols):
        try:
            df = raw[s]
            df = clean_prices(df[["Open", "High", "Low", "Close", "Volume"]])
            if len(df) >= 10:
                out[c] = df
        except Exception:
            pass
    return out


def clean_prices(df):
    """ヤフーのデータの不完全な行を取り除く。
    最新日に出来高0・高値安値0の仮の行が入ることがあり、ATRなどが大きく狂うため。"""
    df = df.copy()
    df = df[(df["Close"] > 0) & df["Close"].notna()]
    df = df[df["Volume"].fillna(0) > 0]
    for col in ("Open", "High", "Low"):
        bad = df[col].isna() | (df[col] <= 0)
        df.loc[bad, col] = df.loc[bad, "Close"]
    df["High"] = df[["High", "Open", "Close"]].max(axis=1)
    df["Low"] = df[["Low", "Open", "Close"]].min(axis=1)
    return df


def save_raw(prices):
    """確認用に、計算に使った直近30日分の株価を CSV で残す"""
    try:
        frames = []
        for c, df in prices.items():
            t = df.tail(30).copy()
            t.insert(0, "Code", c)
            frames.append(t)
        if frames:
            pd.concat(frames).to_csv(os.path.join(HERE, "raw_prices.csv"), encoding="utf-8-sig")
    except Exception:
        pass


def fetch_names(codes, prefill=None):
    import yfinance as yf
    cache = {}
    if os.path.exists(NAME_CACHE):
        try:
            with open(NAME_CACHE, encoding="utf-8") as f:
                cache = json.load(f)
        except Exception:
            cache = {}
    for c, v in (prefill or {}).items():          # 一覧APIの名前・セクターがあれば使う
        if c not in cache and v.get("name"):
            cache[c] = {"name": v.get("name", ""), "sector": v.get("sector", "")}
    todo = [c for c in codes if c not in cache]
    if todo:
        print(f"銘柄名・セクターを取得中（{len(todo)} 銘柄、初回のみ）...")
    for c in todo:
        try:
            info = yf.Ticker(c + ".SA").info
            cache[c] = {"name": info.get("shortName") or info.get("longName") or "",
                        "sector": info.get("sector") or ""}
        except Exception:
            cache[c] = {"name": "", "sector": ""}
        time.sleep(0.3)
    try:
        with open(NAME_CACHE, "w", encoding="utf-8") as f:
            json.dump(cache, f, ensure_ascii=False, indent=1)
    except Exception:
        pass
    return cache


def fetch_earnings(codes):
    """決算予定日（3日以内に調べた銘柄は記録を使う）"""
    import yfinance as yf
    today = dt.date.today()
    cache = {}
    if os.path.exists(EARN_CACHE):
        try:
            with open(EARN_CACHE, encoding="utf-8") as f:
                cache = json.load(f)
        except Exception:
            cache = {}
    res, todo = {}, []
    for c in codes:
        e = cache.get(c)
        if e and (today - dt.date.fromisoformat(e["checked"])).days < 3:
            if e.get("date"):
                d0 = dt.date.fromisoformat(e["date"])
                if d0 >= today:
                    res[c] = d0
        else:
            todo.append(c)
    if todo:
        print(f"決算予定日を取得中（{len(todo)} 銘柄）...")
    for c in todo:
        cache[c] = {"checked": today.isoformat(), "date": None}
        try:
            cal = yf.Ticker(c + ".SA").calendar
            dates = cal.get("Earnings Date") if isinstance(cal, dict) else None
            if dates:
                future = [d for d in dates if isinstance(d, dt.date) and d >= today]
                if future:
                    res[c] = min(future)
                    cache[c]["date"] = res[c].isoformat()
        except Exception:
            pass
        time.sleep(0.2)
    try:
        with open(EARN_CACHE, "w", encoding="utf-8") as f:
            json.dump(cache, f)
    except Exception:
        pass
    return res


# ---------------------------------------------------------------------------
# 指標の計算
# ---------------------------------------------------------------------------
def rsi(close, n=14):
    d = close.diff()
    up = d.clip(lower=0).ewm(alpha=1 / n, adjust=False).mean()
    dn = (-d.clip(upper=0)).ewm(alpha=1 / n, adjust=False).mean()
    rs = up / dn.replace(0, np.nan)
    return 100 - 100 / (1 + rs)


def ma_angle(close, period, look):
    """単純移動平均の傾きを角度(°)にする。
    1本あたりの変化率(%)を x として、角度 = atan(x) を度に直したもの（1本1%で45°）。
    あわせて、終値が移動平均から何%離れているか（乖離率）も返す。"""
    close = close.dropna()
    ma = close.rolling(period).mean()
    if len(ma.dropna()) <= look:
        return None, None
    now, past = ma.iloc[-1], ma.iloc[-1 - look]
    if not past or past <= 0:
        return None, None
    per_bar = ((now / past) ** (1.0 / look) - 1) * 100
    angle = math.degrees(math.atan(per_bar))
    dev = (close.iloc[-1] / now - 1) * 100
    return angle, dev


def resample_close(df, rule):
    """日足の終値を週足・月足の終値に変換（各週・各月の最後の終値）"""
    try:
        return df["Close"].resample(rule).last().dropna()
    except ValueError:            # 古い pandas は月末を "M" と書く
        return df["Close"].resample(rule.replace("ME", "M")).last().dropna()


def compute_metrics(code, df, ma_cfg=None, hourly=None, look=5):
    c, h, l, v = df["Close"], df["High"], df["Low"], df["Volume"]
    prev = c.shift(1)
    tr = pd.concat([h - l, (h - prev).abs(), (l - prev).abs()], axis=1).max(axis=1)
    last = c.iloc[-1]
    ma25 = c.rolling(25).mean()
    value20 = (c * v).rolling(20).mean().iloc[-1]
    vol20 = v.rolling(20).mean()
    rng = ((h - l) / c * 100).iloc[-5:].tolist()
    rng = [np.nan] * (5 - len(rng)) + rng

    def f(x, nd=2):
        try:
            x = float(x)
            return None if math.isnan(x) or math.isinf(x) else round(x, nd)
        except Exception:
            return None

    tr14 = tr.rolling(14).mean().iloc[-1]
    tr5 = tr.iloc[-5:].mean()
    tr20 = tr.rolling(20).mean().iloc[-1]
    logret = np.log(c / prev)
    ma_cfg = ma_cfg or MA_DEFAULTS
    series = {"日足": c, "週足": resample_close(df, "W-FRI"), "月足": resample_close(df, "ME"),
              "時間足": hourly["Close"] if hourly is not None and len(hourly) else pd.Series(dtype=float)}
    ma_vals = {}
    for tf, period, look in ma_cfg:
        a, dv = ma_angle(series[tf], int(period), int(look))
        ma_vals[f"{tf}MA角度(°)"] = f(a, 1)
        ma_vals[f"{tf}MA乖離(%)"] = f(dv)
    # ---- パターン判定用（最新日 n と、その前の look 日）----
    pat = {}
    n = len(c) - 1
    look = max(2, min(int(look or 5), 20))
    if n >= look + 15:
        chg = c.pct_change() * 100
        drops = chg.iloc[n - look:n]
        cpos = n - look + int(np.argmin(drops.values))
        up = 0
        j = n - 1
        while j >= 1 and c.iloc[j] > c.iloc[j - 1]:
            up += 1
            j -= 1
        atr_prev = tr.iloc[n - 14:n].mean()
        rng_today = h.iloc[n] - l.iloc[n]
        pat = {
            "直近の最大下落率(%)": f(drops.min()),
            "急落からの日数": n - cpos,
            "安値更新なし": 1 if l.iloc[n] >= l.iloc[cpos:n].min() else 0,
            "値幅÷ATR": f(rng_today / atr_prev) if atr_prev else None,
            "出来高÷急落日": f(v.iloc[n] / v.iloc[cpos]) if v.iloc[cpos] else None,
            "連続上昇日数": up,
            "5日騰落率・前日まで(%)": f((c.iloc[n - 1] / c.iloc[n - 6] - 1) * 100),
            "高値更新なし": 1 if h.iloc[n] <= h.iloc[n - look:n].max() else 0,
            "上ヒゲ比率": f((h.iloc[n] - c.iloc[n]) / rng_today) if rng_today > 0 else 0,
            "出来高÷直近最大": f(v.iloc[n] / v.iloc[n - look:n].max()) if v.iloc[n - look:n].max() else None,
        }
    return {
        **pat,
        **ma_vals,
        "コード": code,
        "終値(R$)": f(last),
        "前日比(%)": f((last / c.iloc[-2] - 1) * 100),
        "5日騰落率(%)": f((last / c.iloc[-6] - 1) * 100) if len(c) > 6 else None,
        "平均売買代金(百万R$)": f(value20 / 1e6, 1),
        "平均出来高(株)": f(vol20.iloc[-2] if len(vol20) > 1 else vol20.iloc[-1], 0),
        # 最新日を含まない20日平均と比べる
        "出来高倍率(倍)": f(v.iloc[-1] / vol20.iloc[-2]) if vol20.iloc[-2] else None,
        "ATR(%)": f(tr14 / last * 100),
        "ボラ上昇倍率(倍)": f(tr5 / tr20) if tr20 else None,
        "HV20(%)": f(logret.rolling(20).std().iloc[-1] * math.sqrt(252) * 100, 1),
        "25日線乖離(%)": f((last / ma25.iloc[-1] - 1) * 100),
        "RSI(14)": f(rsi(c).iloc[-1], 1),
        "値幅%(4日前)": f(rng[0]), "値幅%(3日前)": f(rng[1]), "値幅%(2日前)": f(rng[2]),
        "値幅%(前日)": f(rng[3]), "値幅%(最新日)": f(rng[4]),
        "最終データ日": df.index[-1].date(),
        "時間足の最終時刻": (hourly.index[-1].strftime("%m/%d %H:%M") + "（現地）") if hourly is not None and len(hourly) else None,
    }


# ---------------------------------------------------------------------------
# Excel 出力
# ---------------------------------------------------------------------------
def read_previous_settings():
    """前回のファイルから設定を読み込む（ユーザーが変えた条件を引き継ぐ）"""
    prev = {"budget": None, "lot": None, "conds": {}, "ma": {}, "pattern": {}, "universe": None, "chart_code": None}
    if not os.path.exists(OUT_FILE):
        return prev
    try:
        wbp = load_workbook(OUT_FILE)
        if "チャート" in wbp.sheetnames:
            prev["chart_code"] = wbp["チャート"]["B1"].value
        ws = wbp["設定"]
        for r in range(1, 80):
            label = ws.cell(r, 1).value
            if not isinstance(label, str):
                continue
            if label == "投資可能額(R$)":       # 旧版の「投資可能額(円)」は単位が違うので引き継がない
                prev["budget"] = ws.cell(r, 2).value
            elif label == "売買単位(株)":
                prev["lot"] = ws.cell(r, 2).value
            elif label == "対象銘柄":
                prev["universe"] = ws.cell(r, 2).value
            elif label in PATTERN_LABELS:
                prev["pattern"][PATTERN_LABELS[label]] = ws.cell(r, 2).value
            elif label in TIMEFRAMES:
                prev["ma"][label] = (ws.cell(r, 2).value, ws.cell(r, 3).value)
            elif r > 8:
                prev["conds"][label] = (ws.cell(r, 2).value, ws.cell(r, 3).value, ws.cell(r, 4).value)
    except Exception as e:
        print("前回の設定は読み込めませんでした（初期値を使います）:", e)
    return prev


def style_header(cell):
    cell.font = Font(name=FONT, bold=True, color="FFFFFF")
    cell.fill = HEADER_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BORDER


def get_ma_config(prev):
    """設定シートで変更された MA の期間・本数を使う（おかしな値なら初期値）"""
    cfg = []
    for tf, period, look in MA_DEFAULTS:
        p, l = prev["ma"].get(tf, (period, look))
        try:
            p, l = int(p), int(l)
            if p < 2 or l < 1 or p > 300 or l > 100:
                raise ValueError
        except Exception:
            p, l = period, look
        cfg.append((tf, p, l))
    return cfg


def pattern_formula(L, r, pat_rows):
    """プリセット（パターン）の判定式。空欄のしきい値はその条件を見ない"""
    P = {k: f"{SET}$B${row}" for k, row in pat_rows.items()}
    X = lambda name: f"{L[name]}{r}"

    def cmp(x, op, p):          # しきい値が空欄なら TRUE
        return f'OR({p}="",AND({x}<>"",{x}{op}{p}))'

    def cmp2(x, op, p):         # しきい値が空欄なら FALSE（OR の中で使う）
        return f'AND({p}<>"",{x}<>"",{x}{op}{p})'

    chg, rsi, hang = X("前日比(%)"), X("RSI(14)"), X("時間足MA角度(°)")
    crash = ",".join([
        cmp(X("直近の最大下落率(%)"), "<=", P["c_drop"]),
        cmp(chg, ">=", P["c_lo"]), cmp(chg, "<=", P["c_hi"]),
        f'{X("安値更新なし")}=1',
        cmp(X("値幅÷ATR"), "<=", P["c_rng"]),
        cmp(X("出来高÷急落日"), "<=", P["c_vol"]),
        cmp(hang, ">=", P["c_hang"]),
        cmp(rsi, ">=", P["c_rsil"]), cmp(rsi, "<=", P["c_rsih"]),
    ])
    rally = ",".join([
        f'OR({cmp2(X("連続上昇日数"), ">=", P["r_days"])},{cmp2(X("5日騰落率・前日まで(%)"), ">=", P["r_pct"])})',
        cmp(chg, "<=", P["r_chg"]),
        f'OR({X("高値更新なし")}=1,{cmp2(X("上ヒゲ比率"), ">=", P["r_wick"])})',
        cmp(X("出来高÷直近最大"), "<=", P["r_vol"]),
        cmp(hang, "<=", P["r_hang"]),
        cmp(rsi, ">=", P["r_rsi"]),
    ])
    sel = P["sel"]
    return f'IF({sel}="{PATTERN_NAMES[1]}",AND({crash}),IF({sel}="{PATTERN_NAMES[2]}",AND({rally}),TRUE))'


def pattern_pass(row, pv, num):
    """pattern_formula と同じ判定を Python で（結果(固定)シート用）"""
    sel = pv.get("sel")
    if sel not in PATTERN_NAMES[1:]:
        return True

    def ok(name, op, key):
        p = num(pv.get(key))
        if p is None:
            return True
        x = num(row.get(name))
        if x is None:
            return False
        return x <= p if op == "<=" else x >= p

    def ok2(name, op, key):
        return num(pv.get(key)) is not None and ok(name, op, key)

    if sel == PATTERN_NAMES[1]:
        return all([ok("直近の最大下落率(%)", "<=", "c_drop"), ok("前日比(%)", ">=", "c_lo"),
                    ok("前日比(%)", "<=", "c_hi"), row.get("安値更新なし") == 1,
                    ok("値幅÷ATR", "<=", "c_rng"), ok("出来高÷急落日", "<=", "c_vol"),
                    ok("時間足MA角度(°)", ">=", "c_hang"), ok("RSI(14)", ">=", "c_rsil"),
                    ok("RSI(14)", "<=", "c_rsih")])
    return all([ok2("連続上昇日数", ">=", "r_days") or ok2("5日騰落率・前日まで(%)", ">=", "r_pct"),
                ok("前日比(%)", "<=", "r_chg"),
                row.get("高値更新なし") == 1 or ok2("上ヒゲ比率", ">=", "r_wick"),
                ok("出来高÷直近最大", "<=", "r_vol"), ok("時間足MA角度(°)", "<=", "r_hang"),
                ok("RSI(14)", ">=", "r_rsi")])


CHART_BARS = {"日足": 150, "週足": 104, "月足": 60}   # チャートに出す本数


def add_chart_sheet(wb, rows, prices, ma_cfg, hits, prev_code=None):
    """「チャート」シート：B1 に銘柄コードを入れると、日足・週足・月足の終値と MA のチャートが切り替わる。
    Numbers でも動くように、式はこのシートの中だけを参照する"""
    if not prices:
        return
    ws = wb.create_sheet("チャート", 1)
    # 対象：売買代金の多い銘柄＋該当銘柄（多すぎると iPad で重くなるため）
    by_code = {r["コード"]: r for r in rows}
    liquid = [r["コード"] for r in sorted(rows, key=lambda r: -(r.get("平均売買代金(百万R$)") or 0))
              if (r.get("平均売買代金(百万R$)") or 0) >= LIQ_FOR_DETAIL]
    codes = list(dict.fromkeys([h["コード"] for h in hits] + liquid))
    codes = [c for c in codes if c in prices][:450]
    if not codes:
        return
    sel = prev_code if prev_code in codes else (hits[0]["コード"] if hits else codes[0])
    per = {tf: p for tf, p, _ in ma_cfg}

    # 足ごとの終値と MA（Python で計算してから、直近の本数だけ使う）
    def series(tf, df):
        c = df["Close"]
        if tf == "週足":
            c = resample_close(df, "W-FRI")
        elif tf == "月足":
            c = resample_close(df, "ME")
        return c, c.rolling(per[tf]).mean()

    BLK = 34                       # 銘柄ごとのデータを置く先頭列（AH）
    LIST = 27                      # 銘柄一覧の先頭列（AA）
    blocks, row0 = {}, 4
    for tf in ("日足", "週足", "月足"):
        n = CHART_BARS[tf]
        data = {c: series(tf, prices[c]) for c in codes}
        idx = sorted(set().union(*[set(v[0].index) for v in data.values()]))[-n:]
        blocks[tf] = (row0, len(idx))
        ws.cell(row0 - 1, BLK - 1, f"{tf}のデータ")
        ws.cell(row0, BLK - 1, "日付")
        for i, d in enumerate(idx):
            ws.cell(row0 + 1 + i, BLK - 1, d.to_pydatetime()).number_format = "yyyy/mm/dd"
        for j, c in enumerate(codes):
            col = BLK + 2 * j
            ws.cell(row0 - 1, col, c)
            ws.cell(row0 - 1, col + 1, c)
            ws.cell(row0, col, "終値")
            ws.cell(row0, col + 1, "MA")
            cl, ma = data[c]
            cl = cl.reindex(idx).ffill(limit=5) if len(cl) else cl
            ma = ma.reindex(idx).ffill(limit=5) if len(ma) else ma
            for i in range(len(idx)):
                v1, v2 = cl.iloc[i], ma.iloc[i]
                if pd.notna(v1):
                    ws.cell(row0 + 1 + i, col, round(float(v1), 2))
                if pd.notna(v2):
                    ws.cell(row0 + 1 + i, col + 1, round(float(v2), 2))
        row0 += len(idx) + 6
    last_col = get_column_letter(BLK + 2 * len(codes) - 1)

    # 銘柄一覧（コード・名前・角度）…入力の候補と、名前・角度の表示に使う
    hdr = ["コード", "銘柄名", "日足MA角度(°)", "週足MA角度(°)", "月足MA角度(°)", "時間足MA角度(°)"]
    for j, h in enumerate(hdr):
        style_header(ws.cell(3, LIST + j, h))
    for i, c in enumerate(codes):
        r = by_code.get(c, {})
        for j, h in enumerate(hdr):
            ws.cell(4 + i, LIST + j, c if h == "コード" else (r.get(h) if h != "銘柄名" else r.get("銘柄名")))
    L1, L2 = get_column_letter(LIST), 4 + len(codes) - 1
    lst = lambda k: f"${get_column_letter(LIST + k)}$4:${get_column_letter(LIST + k)}${L2}"

    # 入力欄と表示
    ws["A1"] = "銘柄コード →"
    ws["A1"].font = Font(name=FONT, bold=True, size=12)
    ws["B1"] = sel
    ws["B1"].fill = INPUT_FILL
    ws["B1"].font = Font(name=FONT, bold=True, size=14, color="0000FF")
    ws["B1"].border = BORDER
    dv = DataValidation(type="list", formula1=f"${L1}$4:${L1}${L2}", allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(ws["B1"])
    m = f"MATCH($B$1,{lst(0)},0)"
    ws["C1"] = f'=IFERROR(INDEX({lst(1)},{m}),"コードが一覧にありません")'
    ws["C1"].font = Font(name=FONT, bold=True, size=12)
    labels = [("E1", "日足角度", 2), ("G1", "週足角度", 3), ("I1", "月足角度", 4), ("K1", "時間足角度", 5)]
    for addr, lab, k in labels:
        col = ws[addr].column
        ws.cell(1, col, lab).font = Font(name=FONT, color="595959")
        cc = ws.cell(1, col + 1, f'=IFERROR(IF(INDEX({lst(k)},{m})="","",INDEX({lst(k)},{m})),"")')
        cc.number_format = '+0.0"°";-0.0"°";0.0"°"'
        cc.font = Font(name=FONT, bold=True)
    ws["A2"] = ("B1 に銘柄コードを入れる（またはリストから選ぶ）と、下の3つのチャートが切り替わります。"
                f"対象は売買代金{LIQ_FOR_DETAIL:g}百万R$以上と該当銘柄の {len(codes)} 銘柄（右側の一覧）。MA の期間は「設定」シートのとおり")
    ws["A2"].font = Font(name=FONT, color="595959")

    # チャート用の表（選んだ銘柄の分だけを、右側のデータから取り出す）
    feeds = {"日足": 1, "週足": 5, "月足": 9}
    anchors = {"日足": "M4", "週足": "M22", "月足": "M40"}
    for tf, fc in feeds.items():
        r0, nrow = blocks[tf]
        ws.cell(3, fc, f"{tf}（MA{per[tf]}）").font = Font(name=FONT, bold=True)
        for j, h in enumerate(["日付", "終値", f"MA{per[tf]}"]):
            style_header(ws.cell(4, fc + j, h))
        rng = f"${get_column_letter(BLK)}${r0 + 1}:${last_col}${r0 + nrow}"
        hdr_rng = f"${get_column_letter(BLK)}${r0 - 1}:${last_col}${r0 - 1}"
        mm = f"MATCH($B$1,{hdr_rng},0)"
        for i in range(nrow):
            rr = 5 + i
            ws.cell(rr, fc, f"={get_column_letter(BLK - 1)}{r0 + 1 + i}").number_format = "yy/mm/dd"
            for k in (0, 1):
                cc = ws.cell(rr, fc + 1 + k,
                             f'=IFERROR(IF(INDEX({rng},{i + 1},{mm}+{k})="",NA(),INDEX({rng},{i + 1},{mm}+{k})),NA())')
                cc.number_format = "0.00"
        ch = LineChart()
        ch.title = f"{tf}　終値と MA{per[tf]}"
        ch.height, ch.width = 8.5, 18
        ch.y_axis.title = "R$"
        ch.legend.position = "b"
        data_ref = Reference(ws, min_col=fc + 1, max_col=fc + 2, min_row=4, max_row=4 + nrow)
        ch.add_data(data_ref, titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=fc, min_row=5, max_row=4 + nrow))
        ch.series[0].graphicalProperties.line.solidFill = "1F3864"
        ch.series[0].graphicalProperties.line.width = 15000
        ch.series[1].graphicalProperties.line.solidFill = "E46C0A"
        ch.series[1].graphicalProperties.line.width = 22000
        for s_ in ch.series:
            s_.smooth = False
        ch.x_axis.number_format = "yy/mm"
        ch.x_axis.delete = False
        ch.y_axis.delete = False
        ws.add_chart(ch, anchors[tf])
    for col in ("A", "E", "I"):
        ws.column_dimensions[col].width = 10
    for col in ("B", "C", "F", "G", "J", "K"):
        ws.column_dimensions[col].width = 9
    ws.column_dimensions["C"].width = 24
    ws.freeze_panes = "A3"


def build_workbook(rows, failed, fetched_at, prev, ma_cfg, prices=None):
    wb = Workbook()

    # ---- 設定シート ----
    s = wb.active
    s.title = "設定"
    s["A1"] = "ブラジル株スクリーナー　設定"
    s["A1"].font = Font(name=FONT, bold=True, size=14)
    s["A2"] = "黄色のセルを変更してください。変えるとすぐ「結果」シートに反映されます。"
    s["A2"].font = Font(name=FONT, color="595959")

    basics = [
        (4, "投資可能額(R$)", prev["budget"] if prev["budget"] is not None else 20000, "#,##0",
         "手持ち資金（レアル）。1銘柄に使える上限額を入れてください", True),
        (5, "売買単位(株)", prev["lot"] if prev["lot"] is not None else 100, "#,##0",
         "B3の通常単位は100株。端株市場（コード末尾F）で1株から買う場合は1", True),
        (6, "データ取得日時", fetched_at, "yyyy/mm/dd hh:mm", "株価は最新の日足（ブラジル時間の終値）", False),
        (7, "対象銘柄", prev.get("universe") if prev.get("universe") in UNIVERSES else UNIVERSES[0], "@",
         "全銘柄＝B3の株式すべて（FII・ETF・BDRは除く）。※変えたら run を実行し直してください", True),
    ]
    dvu = DataValidation(type="list", formula1='"' + ",".join(UNIVERSES) + '"', allow_blank=False)
    s.add_data_validation(dvu)
    for r, label, val, fmt, note, editable in basics:
        s.cell(r, 1, label).font = Font(name=FONT, bold=True)
        c = s.cell(r, 2, val)
        if label == "対象銘柄":
            dvu.add(c)
        c.number_format = fmt
        c.font = Font(name=FONT, color="0000FF" if editable else "000000")
        c.border = BORDER
        if editable:
            c.fill = INPUT_FILL
        s.cell(r, 3, note).font = Font(name=FONT, color="595959")

    for j, hname in enumerate(["移動平均（単純）", "期間(本)", "角度を測る本数", "", "説明"], 1):
        if hname:
            style_header(s.cell(8, j, hname))
    ma_notes = {"時間足": "例：25本（1時間足）を5本分の変化で測る", "日足": "例：25日線を5日分の変化で測る", "週足": "例：13週線を3週分の変化で測る",
                "月足": "例：12か月線を3か月分の変化で測る"}
    for i, (tf, p, l) in enumerate(ma_cfg):
        r = 9 + i
        s.cell(r, 1, tf).font = Font(name=FONT, bold=True)
        for j, val in ((2, p), (3, l)):
            c = s.cell(r, j, val)
            c.fill = INPUT_FILL
            c.font = Font(name=FONT, color="0000FF")
            c.border = BORDER
        s.cell(r, 5, ma_notes[tf] + "。※ここを変えたときは run をもう一度実行すると角度が計算し直されます").font = Font(name=FONT, color="595959")

    PH = 9 + len(ma_cfg) + 1
    for j, hname in enumerate(["プリセット（値動きのパターン）", "値", "", "", "説明"], 1):
        if hname:
            style_header(s.cell(PH, j, hname))
    pat_rows, pat_vals = {}, {}
    pp = prev.get("pattern", {})
    if pp.get("c_hang") == -5 and pp.get("c_rsil") == 25 and pp.get("c_rsih") == 45:   # 初版の初期値のままなら新しい初期値に
        pp.pop("c_hang"); pp.pop("c_rsil"); pp.pop("c_rsih")
    dvp = DataValidation(type="list", formula1='"' + ",".join(PATTERN_NAMES) + '"', allow_blank=False)
    s.add_data_validation(dvp)
    dvm = DataValidation(type="list", formula1='"' + ",".join(PATTERN_MODES) + '"', allow_blank=False)
    s.add_data_validation(dvm)
    for i, (key, label, default, note) in enumerate(PATTERN_ITEMS):
        r = PH + 1 + i
        pat_rows[key] = r
        val = prev["pattern"].get(key, default) if prev.get("pattern") else default
        if key == "sel" and val not in PATTERN_NAMES:
            val = "なし"
        if key == "mode" and val not in PATTERN_MODES:
            val = PATTERN_MODES[0]
        pat_vals[key] = val
        s.cell(r, 1, label).font = Font(name=FONT, bold=True)
        c = s.cell(r, 2, val)
        c.fill = INPUT_FILL
        c.font = Font(name=FONT, color="0000FF", bold=(key == "sel"))
        c.border = BORDER
        c.alignment = Alignment(horizontal="center")
        if key == "sel":
            dvp.add(c)
        if key == "mode":
            dvm.add(c)
        s.cell(r, 5, note).font = Font(name=FONT, color="595959")
    TOP = PH + 1 + len(PATTERN_ITEMS) + 2
    heads = ["条件", "使う", "下限", "上限", "説明"]
    for j, hname in enumerate(heads, 1):
        style_header(s.cell(TOP - 1, j, hname))
    dv = DataValidation(type="list", formula1='"ON,OFF"', allow_blank=False)
    s.add_data_validation(dv)
    cond_rows = {}
    cond_vals = {}
    for i, (key, label, col, on, lo, hi, note) in enumerate(CONDITIONS):
        r = TOP + i
        cond_rows[key] = r
        if label in prev["conds"]:
            on, lo, hi = prev["conds"][label]
        cond_vals[key] = (on, lo, hi)
        s.cell(r, 1, label).font = Font(name=FONT, bold=True)
        c = s.cell(r, 2, on or "OFF")
        dv.add(c)
        c.alignment = Alignment(horizontal="center")
        for j, val in ((2, None), (3, lo), (4, hi)):
            cell = s.cell(r, j) if val is None else s.cell(r, j, val)
            cell.fill = INPUT_FILL
            cell.font = Font(name=FONT, color="0000FF")
            cell.border = BORDER
        if key == "budget":
            for j in (3, 4):
                s.cell(r, j).fill = PatternFill("solid", fgColor="D9D9D9")
                s.cell(r, j).value = None
        s.cell(r, 5, note).font = Font(name=FONT, color="595959")
        s.cell(r, 5).alignment = Alignment(wrap_text=True, vertical="center")
    last_cond = TOP + len(CONDITIONS) - 1
    s.conditional_formatting.add(f"B{TOP}:B{last_cond}", CellIsRule(
        operator="equal", formula=['"ON"'], fill=PatternFill("solid", fgColor="C6EFCE"),
        font=Font(name=FONT, bold=True, color="006100")))
    r = last_cond + 2
    s.cell(r, 1, "・「使う」を ON にした条件だけで絞り込みます（すべて満たす銘柄が残ります）。").font = Font(name=FONT)
    s.cell(r + 1, 1, "・下限・上限は片方だけでも可。空欄ならその側は制限なし。").font = Font(name=FONT)
    s.cell(r + 2, 1, "・%の項目は「3」で3%の意味です。").font = Font(name=FONT)
    if failed:
        s.cell(r + 4, 1, "取得できなかった銘柄（上場廃止・コード変更など）:").font = Font(name=FONT, bold=True)
        s.cell(r + 5, 1, " ".join(failed)).font = Font(name=FONT, color="C00000")
    s.column_dimensions["A"].width = 30
    s.column_dimensions["B"].width = 14
    s.column_dimensions["C"].width = 12
    s.column_dimensions["D"].width = 12
    s.column_dimensions["E"].width = 70
    s.freeze_panes = f"A{TOP}"

    # ---- データシート ----
    d = wb.create_sheet("データ")
    cols = DATA_COLUMNS + ["判定", "連番"]
    L = {name: get_column_letter(i + 1) for i, name in enumerate(cols)}
    for j, name in enumerate(cols, 1):
        style_header(d.cell(1, j, name))
    n = len(rows)
    for i, row in enumerate(rows):
        r = i + 2
        for j, name in enumerate(DATA_COLUMNS, 1):
            if name == "最低購入額(R$)":
                val = f'=IF({L["終値(R$)"]}{r}="","",{L["終値(R$)"]}{r}*{SET}$B$5)'
            else:
                val = row.get(name)
            cell = d.cell(r, j, val)
            cell.font = Font(name=FONT)
        # 判定式
        terms = []
        for key, label, col, *_ in CONDITIONS:
            cr = cond_rows[key]
            on, lo, hi = f"'設定'!$B${cr}", f"'設定'!$C${cr}", f"'設定'!$D${cr}"
            x = f"{L[col]}{r}"
            if key == "budget":
                t = f'IF({on}="ON",AND({x}<>"",{x}<={SET}$B$4),TRUE)'
            elif key == "earn":
                t = f'IF({on}="ON",OR({x}="",AND(OR({lo}="",{x}>={lo}),OR({hi}="",{x}<={hi}))),TRUE)'
            else:
                t = f'IF({on}="ON",AND({x}<>"",OR({lo}="",{x}>={lo}),OR({hi}="",{x}<={hi})),TRUE)'
            if key not in BASE_KEYS:     # プリセット使用中で「投資可能額と流動性だけ」なら、この条件は見ない
                selc, modec = f"{SET}$B${pat_rows['sel']}", f"{SET}$B${pat_rows['mode']}"
                t = f'IF(AND({selc}<>"{PATTERN_NAMES[0]}",{modec}="{PATTERN_MODES[0]}"),TRUE,{t})'
            terms.append(t)
        terms.append(pattern_formula(L, r, pat_rows))
        d.cell(r, len(DATA_COLUMNS) + 1, f'=IF(AND({",".join(terms)}),1,0)')
        jc = d.cell(r, len(DATA_COLUMNS) + 1)
        jc.alignment = Alignment(horizontal="center")
        jc.number_format = '"〇";;""'
        jl, sl = L["判定"], L["連番"]
        d.cell(r, len(DATA_COLUMNS) + 2, f'=IF({jl}{r}=1,SUM(${jl}$2:{jl}{r}),"")')

    fmts = {"急落からの日数": "0", "安値更新なし": "0", "連続上昇日数": "0", "高値更新なし": "0",
            "時間足MA角度(°)": "0.0", "最低購入額(R$)": "#,##0", "日足MA角度(°)": "0.0", "週足MA角度(°)": "0.0", "月足MA角度(°)": "0.0", "平均出来高(株)": "#,##0", "平均売買代金(百万R$)": "#,##0.0",
            "次回決算日": "yyyy/mm/dd", "最終データ日": "yyyy/mm/dd", "決算までの日数": "0"}
    for name, letter in L.items():
        fmt = fmts.get(name, "0.00" if name not in ("コード", "銘柄名", "セクター", "判定", "連番") else None)
        if fmt:
            for r in range(2, n + 2):
                d[f"{letter}{r}"].number_format = fmt
    widths = {"コード": 9, "銘柄名": 22, "セクター": 18}
    for name, letter in L.items():
        d.column_dimensions[letter].width = widths.get(name, 11)
    d.row_dimensions[1].height = 42
    d.freeze_panes = "D2"
    d.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{n + 1}"
    add_color_rules(d, L, 2, n + 1)

    # ---- 結果シート（判定〇の銘柄だけを自動で並べる。並び順は ATR の大きい順）----
    res = wb.create_sheet("結果", 0)
    res["A1"] = "該当数"
    res["A1"].font = Font(name=FONT, bold=True)
    res["B1"] = f'=SUM({DAT}${L["判定"]}$2:${L["判定"]}${n + 1})'
    res["B1"].font = Font(name=FONT, bold=True, size=14, color="C00000")
    res["C1"] = "（ATR% の大きい順。条件は「設定」シートで変更）"
    res["C1"].font = Font(name=FONT, color="595959")
    for j, name in enumerate(DATA_COLUMNS, 1):
        style_header(res.cell(2, j, name))
    sl = L["連番"]
    for k in range(1, n + 1):
        r = k + 2
        for j, name in enumerate(DATA_COLUMNS, 1):
            src = f"'データ'!${L[name]}$2:${L[name]}${n + 1}"
            m = f"MATCH({k},'データ'!${sl}$2:${sl}${n + 1},0)"
            c = res.cell(r, j, f'=IFERROR(IF(INDEX({src},{m})="","",INDEX({src},{m})),"")')
            c.font = Font(name=FONT)
            letter = L[name]
            c.number_format = d[f"{letter}2"].number_format if n else "General"
    for name, letter in L.items():
        if name in DATA_COLUMNS:
            res.column_dimensions[letter].width = d.column_dimensions[letter].width
    res.row_dimensions[2].height = 42
    res.freeze_panes = "D3"
    add_color_rules(res, L, 3, n + 2)
    # ---- 結果（固定）シート：パソコンの設定で絞り込んだ結果を値のまま書く ----
    # Numbers など、式の計算が Excel と違うアプリでも必ず結果が見えるようにするため
    budget = s["B4"].value
    lot = s["B5"].value or 100

    def num(v):
        try:
            return None if v is None or v == "" else float(v)
        except Exception:
            return None

    def passes(row):
        base_only = pat_vals.get("sel") in PATTERN_NAMES[1:] and pat_vals.get("mode") == PATTERN_MODES[0]
        for key, label, col, *_ in CONDITIONS:
            on, lo, hi = cond_vals.get(key, ("OFF", None, None))
            if on != "ON" or (base_only and key not in BASE_KEYS):
                continue
            if key == "budget":
                price = num(row.get("終値(R$)"))
                if price is None or (num(budget) is not None and price * float(lot) > num(budget)):
                    return False
                continue
            x = num(row.get(col))
            if x is None:
                if key == "earn":
                    continue
                return False
            if num(lo) is not None and x < num(lo):
                return False
            if num(hi) is not None and x > num(hi):
                return False
        return pattern_pass(row, pat_vals, num)

    fixed = wb.create_sheet("結果(固定)", 0)
    hits = [r for r in rows if passes(r)]
    fixed["A1"] = f"該当 {len(hits)} 銘柄" + (f"（プリセット：{pat_vals.get('sel')}）" if pat_vals.get("sel") in PATTERN_NAMES[1:] else "")
    fixed["A1"].font = Font(name=FONT, bold=True, size=14, color="C00000")
    fixed["C1"] = f"パソコンの「設定」シートの条件で絞り込んだ結果（{fetched_at:%Y/%m/%d %H:%M} 時点）。iPad で条件を変えても、このシートは変わりません"
    fixed["C1"].font = Font(name=FONT, color="595959")
    for j, name in enumerate(DATA_COLUMNS, 1):
        style_header(fixed.cell(2, j, name))
    for i, row in enumerate(hits):
        for j, name in enumerate(DATA_COLUMNS, 1):
            if name == "最低購入額(R$)":
                p_ = num(row.get("終値(R$)"))
                val = round(p_ * float(lot)) if p_ is not None else None
            else:
                val = row.get(name)
            c = fixed.cell(i + 3, j, val)
            c.font = Font(name=FONT)
            c.number_format = d[f"{L[name]}2"].number_format if n else "General"
    for name, letter in L.items():
        if name in DATA_COLUMNS:
            fixed.column_dimensions[letter].width = d.column_dimensions[letter].width
    fixed.row_dimensions[2].height = 42
    fixed.freeze_panes = "C3"
    add_color_rules(fixed, L, 3, max(3, len(hits) + 2))

    try:
        add_chart_sheet(wb, rows, prices, ma_cfg, hits, prev.get("chart_code"))
    except Exception as e:
        print("チャートシートを作れませんでした:", e)
    wb.active = 0
    return wb


def add_color_rules(ws, L, r1, r2):
    green = Font(name=FONT, color="007A33")
    red = Font(name=FONT, color="C00000")
    for col in ("前日比(%)", "5日騰落率(%)"):
        rng = f"{L[col]}{r1}:{L[col]}{r2}"
        ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["0"], font=green))
        ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"], font=red))
    rng = f"{L['ボラ上昇倍率(倍)']}{r1}:{L['ボラ上昇倍率(倍)']}{r2}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=["1.3"],
                                                  fill=PatternFill("solid", fgColor="FCE4D6")))
    rng = f"{L['出来高倍率(倍)']}{r1}:{L['出来高倍率(倍)']}{r2}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=["2"],
                                                  fill=PatternFill("solid", fgColor="FCE4D6")))


# ---------------------------------------------------------------------------
def demo_hourly(codes):
    rng = np.random.default_rng(2)
    idx = pd.date_range(end=pd.Timestamp.now().floor("h"), periods=300, freq="h")
    out = {}
    for c in codes:
        close = rng.uniform(3, 80) * np.exp(np.cumsum(rng.normal(0, 0.006, len(idx))))
        out[c] = pd.DataFrame({"Open": close, "High": close * 1.003, "Low": close * 0.997,
                               "Close": close, "Volume": 1e5}, index=idx)
    return out


def demo_data(codes):
    """通信なしで動作確認するためのダミーデータ（--demo のときだけ使用）"""
    rng = np.random.default_rng(1)
    out = {}
    idx = pd.bdate_range(end=dt.date.today(), periods=250)
    for c in codes:
        p0 = rng.uniform(3, 80)
        vol = rng.uniform(0.01, 0.04)
        ret = rng.normal(0.0003, vol, len(idx))
        close = p0 * np.exp(np.cumsum(ret))
        hi = close * (1 + np.abs(rng.normal(0, vol / 2, len(idx))))
        lo = close * (1 - np.abs(rng.normal(0, vol / 2, len(idx))))
        v = rng.uniform(2e5, 3e7) * rng.lognormal(0, 0.4, len(idx))
        out[c] = pd.DataFrame({"Open": close, "High": hi, "Low": lo, "Close": close, "Volume": v}, index=idx)
    return out


def copy_to_cloud(src):
    """copy_to.txt に書かれたフォルダ（OneDrive など）に結果をコピーする。iPad で見るため。"""
    cfg = os.path.join(HERE, "copy_to.txt")
    if not os.path.exists(cfg):
        return
    import shutil
    try:
        with open(cfg, encoding="utf-8-sig") as f:
            dest = f.readline().strip().strip('"')
        if not dest:
            return
        os.makedirs(dest, exist_ok=True)
        shutil.copy2(src, os.path.join(dest, os.path.basename(OUT_FILE)))
        print("コピーしました →", dest)
    except Exception as e:
        print("コピーできませんでした:", e)


def main():
    demo = "--demo" in sys.argv
    prev = read_previous_settings()
    universe = prev.get("universe") if prev.get("universe") in UNIVERSES else UNIVERSES[0]
    if demo:
        codes, tinfo = list(dict.fromkeys(DEFAULT_TICKERS)), {}
    else:
        codes, tinfo = load_tickers(universe)
    if demo:
        prices = demo_data(codes)
        hourly = demo_hourly(codes)
        names = {c: {"name": f"(ダミー){c}", "sector": "Demo"} for c in codes}
        earnings = {codes[0]: dt.date.today() + dt.timedelta(days=3)}
    else:
        prices = fetch_prices(codes)
        if not prices:
            print("株価を1銘柄も取得できませんでした。ネット接続か、ヤフー側の仕様変更の可能性があります。")
            print("「python -m pip install -U yfinance」で yfinance を更新すると直ることがあります。")
            if "--auto" not in sys.argv:
                input("Enterキーで終了")
            return
        save_raw(prices)
        # 売買代金の少ない銘柄は、1時間足と決算日の取得を省いて時間を短くする
        liquid = [c for c, df in prices.items()
                  if (df["Close"] * df["Volume"]).tail(20).mean() / 1e6 >= LIQ_FOR_DETAIL]
        print(f"売買代金 {LIQ_FOR_DETAIL:g}百万R$以上: {len(liquid)} 銘柄（1時間足・決算日はこの銘柄だけ取得）")
        hourly = fetch_hourly(liquid)
        names = fetch_names(list(prices), tinfo)
        earnings = fetch_earnings(liquid) if FETCH_EARNINGS else {}

    ma_cfg = get_ma_config(prev)
    try:
        pat_look = int(prev["pattern"].get("look", 5))
    except Exception:
        pat_look = 5
    failed = [c for c in codes if c not in prices]
    today = dt.date.today()
    rows = []
    for c, df in prices.items():
        try:
            m = compute_metrics(c, df, ma_cfg, hourly.get(c), pat_look)
        except Exception as e:
            print(f"{c}: 計算できませんでした ({e})")
            failed.append(c)
            continue
        m["銘柄名"] = names.get(c, {}).get("name", "")
        m["セクター"] = names.get(c, {}).get("sector", "")
        e = earnings.get(c)
        m["次回決算日"] = e
        m["決算までの日数"] = (e - today).days if e else None
        rows.append(m)
    rows.sort(key=lambda m: -(m.get("ATR(%)") or 0))

    wb = build_workbook(rows, failed, dt.datetime.now().replace(microsecond=0), prev, ma_cfg, prices)
    out = OUT_FILE
    try:
        wb.save(out)
    except PermissionError:
        out = OUT_FILE.replace(".xlsx", "_最新.xlsx")
        wb.save(out)
        print("※ 元のファイルが Excel で開かれていたため、「_最新」の名前で保存しました（OneDrive には通常の名前でコピーします）。")
    print(f"完了: {len(rows)} 銘柄 → {out}")
    copy_to_cloud(out)
    if failed:
        print("取得できなかった銘柄:", " ".join(failed))


if __name__ == "__main__":
    main()
